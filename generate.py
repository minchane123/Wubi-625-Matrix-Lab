# -*- coding: utf-8 -*-
import os
import pandas as pd
import string
from openpyxl.styles import PatternFill, Alignment, Border, Side, Font
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.utils import get_column_letter

def load_freq_and_level(file_path):
    char_rank = {}
    tier_2500 = set()
    tier_3500 = set()
    tier_5000 = set()
    if not os.path.exists(file_path):
        return char_rank, tier_2500, tier_3500, tier_5000

    df = pd.read_csv(file_path)
    for _, row in df.iterrows():
        char = row['character']
        try:
            rank = int(row['frequency_rank'])
        except:
            rank = 9999999
        char_rank[char] = rank
        if 1 <= rank <= 2500:
            tier_2500.add(char)
        elif 2500 < rank <= 3500:
            tier_3500.add(char)
        elif 3500 < rank <= 5000:
            tier_5000.add(char)
    return char_rank, tier_2500, tier_3500, tier_5000

def export_to_pdf(excel_path, pdf_path):
    try:
        import win32com.client
        print("🖨️ 正在调用本地 Excel 导出 PDF...")
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        
        wb = excel.Workbooks.Open(excel_path)
        wb.Worksheets(['二级简码矩阵', '词频同前缀候选']).Select()
        wb.ActiveSheet.ExportAsFixedFormat(0, pdf_path)
        
        wb.Close(False)
        excel.Quit()
        print(f"🎉 PDF 导出成功：\n👉 {pdf_path}")
    except Exception as e:
        print(f"⚠️ PDF 导出失败，请检查依赖或文件占用。错误信息：{e}")

def generate_matrix():
    current_dir = os.path.dirname(os.path.abspath(__file__))
    data_dir = os.path.join(current_dir, "data")
    output_dir = os.path.join(current_dir, "output")
    os.makedirs(data_dir, exist_ok=True)
    os.makedirs(output_dir, exist_ok=True)

    csv_path = os.path.join(data_dir, "系统码表.csv")
    excel_path = os.path.join(output_dir, "五笔625炼药台.xlsx")
    pdf_path = os.path.join(output_dir, "五笔625炼药台.pdf")
    hanzidb_path = os.path.join(data_dir, "hanzi_db.csv")

    char_rank, TIER_2500, TIER_3500, TIER_5000 = load_freq_and_level(hanzidb_path)

    if not os.path.exists(csv_path): return

    dict_2char = {}  
    dict_prefix = {} 
    encodings = ['utf-8', 'gbk', 'utf-8-sig']
    for enc in encodings:
        try:
            with open(csv_path, 'r', encoding=enc) as f: lines = f.readlines()
            break
        except UnicodeDecodeError: continue

    valid_letters = set(string.ascii_lowercase[:25])
    line_idx = 0
    for line in lines:
        line_idx += 1
        line = line.strip()
        if not line or line.startswith('#'): continue
        parts = line.replace('\t', ',').replace(' ', ',').split(',')
        parts = [p.strip() for p in parts if p.strip()]
        if len(parts) < 2: continue
        code = parts[0].lower()
        if 2 <= len(code) <= 4 and all(c in valid_letters for c in code):
            single_chars = [word for word in parts[1:] if len(word) == 1]
            if not single_chars: continue
            prefix = code[:2]
            suffix = code[2:] 
            if len(code) == 2:
                if code not in dict_2char: dict_2char[code] = single_chars[0]
            elif len(code) in (3, 4):
                if prefix not in dict_prefix: dict_prefix[prefix] = []
                for char in single_chars:
                    if not any(item[0] == char for item in dict_prefix[prefix]):
                        dict_prefix[prefix].append((char, suffix, line_idx))

    sheet1_chars_set = set(dict_2char.values())
    final_prefix_data = {}
    for prefix, items in dict_prefix.items():
        valid_items = [(char, suffix, char_rank.get(char, 9999999)) for char, suffix, line_num in items if char_rank.get(char, 9999999) <= 6000]
        valid_items.sort(key=lambda x: x[2])
        final_prefix_data[prefix] = [(char, "" if char in sheet1_chars_set else suffix) for char, suffix, _ in valid_items]

    letters = [c for c in string.ascii_lowercase if c != 'z']
    level_1_simple = {'a': '工', 'b': '了', 'c': '以', 'd': '在', 'e': '有', 'f': '地', 'g': '一', 'h': '上', 'i': '不', 'j': '是', 'k': '中', 'l': '国', 'm': '同', 'n': '民', 'o': '为', 'p': '这', 'q': '我', 'r': '的', 's': '要', 't': '和', 'u': '产', 'v': '发', 'w': '人', 'x': '经', 'y': '主'}

    matrix_data = []
    for first_char in letters:
        row = {"首码": first_char.upper()}
        char_l1 = level_1_simple.get(first_char, "")
        if char_l1:
            rank = char_rank.get(char_l1, 9999999)
            row["一级简码"] = f"{char_l1}|{str(rank).zfill(4) if rank < 9999999 else '----'}"
        else:
            row["一级简码"] = ""

        for second_char in letters:
            code = first_char + second_char
            char = dict_2char.get(code, "")
            if char:
                rank = char_rank.get(char, 9999999)
                row[second_char.upper()] = f"{char}|{str(rank).zfill(4) if rank < 9999999 else '----'}"
            else:
                row[second_char.upper()] = ""
        matrix_data.append(row)

    df_2char = pd.DataFrame(matrix_data)[['首码', '一级简码'] + [c.upper() for c in letters]]

    list_prefix_rows = []
    block_ranges = [] 
    current_excel_row = 2 
    for first_char in letters:
        for second_char in letters:
            prefix = first_char + second_char
            chars = final_prefix_data.get(prefix, [])
            
            if not chars:
                row = {"前缀编码": prefix.upper()}
                for i in range(10): row[f"候选{i+1}"] = ""
                list_prefix_rows.append(row)
                block_ranges.append((current_excel_row, current_excel_row, prefix.upper()))
                current_excel_row += 1
            else:
                chunks = [chars[i:i + 10] for i in range(0, len(chars), 10)]
                start_row = current_excel_row
                for chunk in chunks:
                    row = {"前缀编码": prefix.upper()}
                    for i in range(10):
                        if i < len(chunk):
                            char, suffix = chunk[i]
                            rank = char_rank.get(char, 9999999)
                            row[f"候选{i+1}"] = f"{char}{suffix}|{str(rank).zfill(4) if rank < 9999999 else '----'}"
                        else:
                            row[f"候选{i+1}"] = ""
                    list_prefix_rows.append(row)
                    current_excel_row += 1
                end_row = current_excel_row - 1
                block_ranges.append((start_row, end_row, prefix.upper()))

    df_prefix_all = pd.DataFrame(list_prefix_rows)

    try:
        with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
            df_2char.to_excel(writer, sheet_name='二级简码矩阵', index=False)
            df_prefix_all.to_excel(writer, sheet_name='词频同前缀候选', index=False)

            workbook = writer.book
            green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            blue_fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
            orange_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

            left_align = Alignment(horizontal='left', vertical='center', indent=0)
            code_align = Alignment(horizontal='center', vertical='top')

            font_char = InlineFont(rFont='新宋体', sz=11)
            font_rank = InlineFont(rFont='新宋体', sz=8, color='808080')

            def apply_rich_text_and_style(cell):
                if not cell.value: return
                val_str = str(cell.value)
                if '|' not in val_str: return

                char_part, rank_part = val_str.split('|')
                visual_width = sum(2 if '\u4e00' <= c <= '\u9fa5' else 1 for c in char_part)
                spaces = " " * max(0, 4 - visual_width)
                
                cell.value = CellRichText(
                    TextBlock(font_char, f"{char_part}{spaces}"),
                    TextBlock(font_rank, rank_part)
                )

                char = char_part[0]
                if char in TIER_2500: cell.fill = green_fill
                elif char in TIER_3500: cell.fill = blue_fill
                elif char in TIER_5000: cell.fill = orange_fill

                cell.alignment = left_align

            # 处理 Sheet1
            ws1 = workbook['二级简码矩阵']
            ws1.column_dimensions['A'].width = 10
            ws1.column_dimensions['B'].width = 12
            for col_letter in letters:
                ws1.column_dimensions[col_letter.upper()].width = 11
            
            for row in ws1.iter_rows(min_row=1, max_row=ws1.max_row, min_col=1, max_col=1):
                for cell in row:
                    cell.font = cell.font.copy(bold=True)
            for row in ws1.iter_rows(min_row=2, min_col=2, max_col=len(df_2char.columns)):
                for cell in row:
                    apply_rich_text_and_style(cell)

            # 处理 Sheet2
            ws2 = workbook['词频同前缀候选']
            ws2.column_dimensions['A'].width = 12
            for col_idx in range(2, 12):
                col_letter = get_column_letter(col_idx)
                ws2.column_dimensions[col_letter].width = 12

            for row in ws2.iter_rows(min_row=2, min_col=2, max_col=11):
                for cell in row:
                    apply_rich_text_and_style(cell)

            # ================= 终极线宽修复 =================
            # 使用 Excel 里最细的 hair 线型
            hair_side = Side(style='hair', color='D9D9D9') 
            
            for start_row, end_row, prefix_code in block_ranges:
                for r in range(start_row, end_row + 1):
                    ws2.cell(row=r, column=1).value = ""
                    
                    # 取消相邻两块的边框重叠：只在块的最后一行画底线，并强制设为 hair 极细线
                    if r == end_row:
                        for c in range(1, 12):
                            cell = ws2.cell(row=r, column=c)
                            # 保持原本其他边框，只覆盖 bottom
                            current_border = cell.border
                            cell.border = Border(
                                top=current_border.top,
                                bottom=hair_side,
                                left=current_border.left,
                                right=current_border.right
                            )
                
                top_cell = ws2.cell(row=start_row, column=1)
                top_cell.value = prefix_code
                top_cell.font = Font(name='新宋体', bold=True, size=11)
                top_cell.alignment = code_align 
            # =================================================

            # ================= 彻底解决 A4 排版 =================
            # 关闭默认的浅色网格线打印，只显示我们画的线
            ws1.print_options.gridLines = False
            ws2.print_options.gridLines = False

            # Sheet1: 强制 A4 横向，通过 fitToWidth = 2 让 Excel 自动切成两页宽，不挤压
            ws1.page_setup.paperSize = ws1.PAPERSIZE_A4
            ws1.page_setup.orientation = ws1.ORIENTATION_LANDSCAPE
            ws1.sheet_properties.pageSetUpPr.fitToPage = True
            ws1.page_setup.fitToWidth = 2   # 强制分成 2 页宽
            ws1.page_setup.fitToHeight = 0  # 高度无限
            ws1.page_margins.left = 0.2
            ws1.page_margins.right = 0.2
            ws1.page_margins.top = 0.3
            ws1.page_margins.bottom = 0.3
            ws1.print_title_rows = '1:1' 
            ws1.print_title_cols = 'A:B' # 每页都带出前两列
            
            # Sheet2: 强制 A4 横向，缩放至 1 页宽
            ws2.page_setup.paperSize = ws2.PAPERSIZE_A4
            ws2.page_setup.orientation = ws2.ORIENTATION_LANDSCAPE
            ws2.sheet_properties.pageSetUpPr.fitToPage = True
            ws2.page_setup.fitToWidth = 1
            ws2.page_setup.fitToHeight = 0
            ws2.page_margins.left = 0.2
            ws2.page_margins.right = 0.2
            ws2.page_margins.top = 0.3
            ws2.page_margins.bottom = 0.3
            ws2.print_title_rows = '1:1'
            # ====================================================

        print(f"🎉 完美！Excel 已生成：\n👉 {excel_path}")
        export_to_pdf(excel_path, pdf_path)

    except PermissionError:
        print(f"❌ 导出失败：请先关闭已经打开的 {excel_path} 文件！")
    except Exception as e:
        print(f"❌ 发生错误：{e}")

if __name__ == "__main__":
    generate_matrix()